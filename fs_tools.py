"""Structured filesystem tools. Relative paths are resolved inside FS_ROOT.

Set FS_ROOT before importing this module, or use the project directory default.
These tools are intended for a local, single-user educational application.
"""
import os
import re
import tempfile
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(os.getenv('FS_ROOT', Path(__file__).parent)).resolve()
MAX_BYTES = 10 * 1024 * 1024


def _path(value: str) -> Path:
    if not isinstance(value, str) or not value.strip():
        raise ValueError('Path must be a non-empty string.')
    path = (ROOT / value).resolve()
    if not path.is_relative_to(ROOT):
        raise ValueError('Path is outside the allowed workspace.')
    if any(part.startswith('.') for part in path.relative_to(ROOT).parts):
        raise ValueError('Hidden files and directories are not accessible.')
    return path


def _metadata(path: Path) -> dict:
    stat = path.stat()
    return {'name': path.name, 'filepath': str(path.relative_to(ROOT)),
            'extension': path.suffix.lower(), 'size': stat.st_size,
            'modified_date': datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat()}


def _error(exc: Exception) -> dict:
    return {'success': False, 'error': type(exc).__name__, 'message': str(exc)}


def read_file(filepath: str) -> dict:
    """Extract UTF-8 TXT, text-based PDF, or DOCX text and metadata."""
    try:
        path = _path(filepath)
        if not path.is_file():
            raise FileNotFoundError(f'File not found: {filepath}')
        if path.stat().st_size > MAX_BYTES:
            raise ValueError('File exceeds the 10 MiB limit.')
        metadata = _metadata(path)
        if path.suffix.lower() == '.txt':
            content = path.read_text(encoding='utf-8-sig')
        elif path.suffix.lower() == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(path)
            if reader.is_encrypted:
                raise ValueError('Encrypted PDFs are not supported.')
            content = '\n'.join(page.extract_text() or '' for page in reader.pages)
            metadata['pages'] = len(reader.pages)
        elif path.suffix.lower() == '.docx':
            from docx import Document
            document = Document(path)
            parts = [paragraph.text for paragraph in document.paragraphs]
            parts.extend(' | '.join(cell.text for cell in row.cells)
                         for table in document.tables for row in table.rows)
            content = '\n'.join(parts)
        else:
            raise ValueError('Supported read formats: .pdf, .txt, .docx')
        metadata['characters'] = len(content)
        return {'success': True, 'content': content, 'metadata': metadata,
                'warning': None if content.strip() else 'No text extracted. Scanned PDFs require OCR.'}
    except Exception as exc:
        return _error(exc)


def list_files(directory: str, extension: str = None) -> list:
    """List immediate files in name order; an empty directory returns [].

    Errors return a one-element list containing a structured error object,
    preserving the assignment's list return type. Hidden files are excluded.
    """
    try:
        path = _path(directory)
        if not path.is_dir():
            raise NotADirectoryError(f'Directory not found: {directory}')
        if extension is not None:
            if not isinstance(extension, str) or not extension.strip():
                raise ValueError('Extension must be a non-empty string or null.')
            extension = '.' + extension.lower().lstrip('.')
        result = []
        for item in sorted(path.iterdir(), key=lambda p: p.name.lower()):
            if item.name.startswith('.') or not item.is_file():
                continue
            safe = _path(str(item))
            if extension is None or safe.suffix.lower() == extension:
                result.append(_metadata(safe))
        return result
    except Exception as exc:
        return [_error(exc)]


def write_file(filepath: str, content: str) -> dict:
    """Atomically write UTF-8 text; create parents. Do not write PDF/DOCX binaries."""
    temporary = None
    try:
        path = _path(filepath)
        if path.suffix.lower() not in {'.txt', '.md'}:
            raise ValueError('Write targets must end in .txt or .md.')
        if not isinstance(content, str):
            raise ValueError('Content must be a string.')
        if len(content.encode('utf-8')) > MAX_BYTES:
            raise ValueError('Content exceeds the 10 MiB limit.')
        path.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(mode='w', encoding='utf-8', dir=path.parent,
                                         delete=False) as output:
            temporary = Path(output.name)
            output.write(content)
        temporary.replace(path)
        return {'success': True, 'message': 'File written successfully.', 'metadata': _metadata(path)}
    except Exception as exc:
        return _error(exc)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def search_in_file(filepath: str, keyword: str) -> dict:
    """Find literal, case-insensitive matches with 70 characters of context."""
    if not isinstance(keyword, str) or not keyword.strip():
        return _error(ValueError('Keyword must be a non-empty string.'))
    result = read_file(filepath)
    if not result['success']:
        return result
    content = result['content']
    matches = [{'text': match.group(), 'start': match.start(), 'end': match.end(),
                'line': content.count('\n', 0, match.start()) + 1,
                'context': content[max(0, match.start()-70):match.end()+70]}
               for match in re.finditer(re.escape(keyword), content, re.IGNORECASE)]
    return {'success': True, 'filepath': filepath, 'keyword': keyword,
            'count': len(matches), 'matches': matches}
